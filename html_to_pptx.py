"""
html_to_pptx.py

Parses a saved HTML page (headings, paragraphs, images) and reassembles the
content into a PowerPoint slide deck -- one slide per heading section, with
that section's body text and image placed automatically.

Useful for converting static HTML exports (e.g., an old photo-book or
scrapbook page export) into a shareable, presentable format.

Usage:
    python html_to_pptx.py --html scrapbook_page.html --out slideshow.pptx --base-url https://example.com
"""

import argparse
from pathlib import Path
from urllib.parse import urljoin, urlparse

import requests
from bs4 import BeautifulSoup
from pptx import Presentation
from pptx.util import Inches, Pt


def download_image(img_url: str, output_dir: Path):
    output_dir.mkdir(parents=True, exist_ok=True)
    try:
        response = requests.get(img_url, timeout=20)
        response.raise_for_status()

        name = Path(urlparse(img_url).path).name or "image.jpg"
        if "." not in name:
            name += ".jpg"

        img_path = output_dir / name
        img_path.write_bytes(response.content)
        return img_path
    except Exception:
        return None


def add_textbox(slide, left, top, width, height, text, font_size=18, bold=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    text_frame = box.text_frame
    text_frame.clear()
    paragraph = text_frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.size = Pt(font_size)
    paragraph.font.bold = bold
    return box


def html_to_pptx(html_path: str, pptx_path: str, base_url: str = None, image_dir: str = "downloaded_images"):
    html_path = Path(html_path)
    html_content = html_path.read_text(encoding="utf-8", errors="ignore")
    soup = BeautifulSoup(html_content, "lxml")

    presentation = Presentation()
    title_layout = presentation.slide_layouts[0]
    blank_layout = presentation.slide_layouts[6]

    doc_title = soup.title.string.strip() if soup.title and soup.title.string else html_path.stem
    title_slide = presentation.slides.add_slide(title_layout)
    title_slide.shapes.title.text = doc_title

    current_title, current_body, current_image = None, [], None

    def flush_slide():
        nonlocal current_title, current_body, current_image
        if not (current_title or current_body or current_image):
            return

        slide = presentation.slides.add_slide(blank_layout)

        if current_title:
            add_textbox(slide, Inches(0.6), Inches(0.4), Inches(8.5), Inches(0.8),
                        current_title, font_size=26, bold=True)

        if current_body:
            body_text = "\n".join(current_body[:8])
            add_textbox(slide, Inches(0.8), Inches(1.4), Inches(4.6), Inches(4.8),
                        body_text, font_size=18)

        if current_image and Path(current_image).exists():
            slide.shapes.add_picture(str(current_image), Inches(5.5), Inches(1.5), width=Inches(3.5))

        current_title, current_body, current_image = None, [], None

    for tag in soup.find_all(["h1", "h2", "p", "img"]):
        if tag.name in ("h1", "h2"):
            flush_slide()
            current_title = tag.get_text(" ", strip=True)
        elif tag.name == "p":
            text = tag.get_text(" ", strip=True)
            if text:
                current_body.append(text)
        elif tag.name == "img" and current_image is None:
            src = tag.get("src")
            if not src:
                continue
            img_url = urljoin(base_url, src) if base_url else src
            downloaded = download_image(img_url, Path(image_dir))
            if downloaded:
                current_image = downloaded

    flush_slide()
    presentation.save(pptx_path)
    print(f"Saved presentation to {pptx_path}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Convert a saved HTML page into a PowerPoint slide deck.")
    parser.add_argument("--html", required=True, help="Path to the source HTML file")
    parser.add_argument("--out", required=True, help="Output .pptx path")
    parser.add_argument("--base-url", default=None, help="Base URL for resolving relative image paths")
    parser.add_argument("--image-dir", default="downloaded_images", help="Directory to save downloaded images")
    args = parser.parse_args()

    html_to_pptx(args.html, args.out, args.base_url, args.image_dir)
